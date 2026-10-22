"""Generate roll-number-only QR codes from student metadata.

Install: python -m pip install -r scripts/requirements-student-qr.txt
Run: python scripts/generate_student_qr.py
Optional: --input path/to/workbook.xlsx --output path/to/student_qr

Only rollno and name are read. All rows with a duplicated roll number are
skipped, rather than choosing an ambiguous student's name. Blank rows are
ignored; invalid rows are reported. The source workbook is opened read-only.
"""
import argparse
import csv
import os
import re
import sys
import tempfile
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_students(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = [cell_text(value).lower() for value in next(rows, ())]
        for required in ("rollno", "name"):
            if header.count(required) != 1:
                raise ValueError(f"Expected exactly one '{required}' column in sheet '{sheet.title}'.")
        roll_index, name_index = header.index("rollno"), header.index("name")
        records = []
        for row_number, values in enumerate(rows, start=2):
            roll = cell_text(values[roll_index] if roll_index < len(values) else None)
            name = cell_text(values[name_index] if name_index < len(values) else None)
            if not roll and not name:
                continue
            records.append({"row": row_number, "rollno": roll, "name": name})
        return records
    finally:
        workbook.close()


def generate(input_path, output_dir):
    import qrcode
    records = read_students(input_path)
    counts = Counter(record["rollno"] for record in records if record["rollno"])
    filename_counts = Counter(record["rollno"].casefold() for record in records if record["rollno"])
    duplicates = {roll: count for roll, count in counts.items() if count > 1}
    valid, skipped = [], []
    for record in records:
        roll = record["rollno"]
        reason = None
        if not roll:
            reason = "Empty roll number"
        elif roll in duplicates:
            reason = "Duplicate roll number (all conflicting rows skipped)"
        elif filename_counts[roll.casefold()] > 1:
            reason = "Roll numbers collide as filenames on Windows (all conflicting rows skipped)"
        elif (re.search(r'[<>:"/\\|?*\x00-\x1f]', roll) or roll in (".", "..")
              or roll.endswith((".", " ")) or len(roll) > 180
              or re.fullmatch(r"(?i)(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?", roll)):
            reason = "Roll number is unsafe for a PNG filename"
        if reason:
            skipped.append({**record, "reason": reason})
        else:
            valid.append(record)

    output_dir.mkdir(parents=True, exist_ok=True)
    # Stage every output before replacing existing files. Rerunning overwrites
    # the same filenames and manifest without touching the input or photos.
    with tempfile.TemporaryDirectory(prefix=".qr-stage-", dir=output_dir) as staging:
        stage = Path(staging)
        for record in valid:
            qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=10, border=4)
            qr.add_data(record["rollno"])
            qr.make(fit=True)
            qr.make_image(fill_color="black", back_color="white").save(stage / f"{record['rollno']}.png")
        with (stage / "student_qr_list.csv").open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["rollno", "name", "qr_file"])
            writer.writeheader()
            for record in valid:
                writer.writerow({"rollno": record["rollno"], "name": record["name"],
                                 "qr_file": (output_dir / f"{record['rollno']}.png").relative_to(output_dir.parent).as_posix()})
        for file in stage.iterdir():
            os.replace(file, output_dir / file.name)

    report = {"total_students": len(records), "qr_codes_generated": len(valid),
              "skipped_records": skipped, "duplicates": duplicates}
    print(f"Total students: {len(records)}")
    print(f"QR codes generated: {len(valid)}")
    print(f"Skipped records: {len(skipped)}")
    print(f"Duplicates: {len(duplicates)} roll numbers ({sum(duplicates.values())} conflicting rows)")
    for roll, count in duplicates.items():
        print(f"  Duplicate {roll}: {count} rows")
    for record in skipped:
        print(f"  Row {record['row']}: {record['rollno'] or '<empty>'} — {record['reason']}")
    print(f"Manifest: {output_dir / 'student_qr_list.csv'}")
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", type=Path, default=ROOT / "student_face_list.xlsx")
    parser.add_argument("--output", type=Path, default=ROOT / "student_qr")
    args = parser.parse_args()
    try:
        report = generate(args.input.resolve(), args.output.resolve())
    except ImportError as error:
        print(f"Missing dependency: {error.name}. Install scripts/requirements-student-qr.txt", file=sys.stderr)
        return 1
    except (OSError, ValueError) as error:
        print(f"QR generation failed: {error}", file=sys.stderr)
        return 1
    return 2 if report["skipped_records"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
