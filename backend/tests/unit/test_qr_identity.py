from pathlib import Path
from unittest.mock import Mock

import cv2
import numpy as np
import pytest
from openpyxl import Workbook

from ai_engine.identity.qr_identity import QRIdentityTracker, load_student_mapping
from ai_engine.vision.frame_annotator import FrameAnnotator


def person(tid=12, bbox=(20,20,250,350), **extra):
    return dict(track_id=tid,class_name="person",bbox=list(bbox),confidence=.91,**extra)


def corners(x,y):
    return np.array([[x-10,y-10],[x+10,y-10],[x+10,y+10],[x-10,y+10]],dtype=float)


def scanner(values=("S123",), points=None, **kwargs):
    tracker = QRIdentityTracker({"S123":"Example Name"},**kwargs)
    tracker.decoder = Mock()
    tracker.decoder.detectAndDecode.return_value = ("",None,None)
    tracker.decoder.detectAndDecodeMulti.return_value = (True,values,np.array(points if points is not None else [corners(100,200)]),None)
    tracker.decoder.detectAndDecodeMulti.side_effect = lambda image: tracker.decoder.detectAndDecodeMulti.return_value if image.shape[:2] == (400,600) else (False,(),None,None)
    return tracker


def test_mapping_caches_profile_metadata_without_reading_photos(tmp_path):
    path=tmp_path/'students.xlsx'
    workbook=Workbook();workbook.active.append(['rollno','name','face_file'])
    workbook.active.append(['S123','Example Name','does-not-exist.jpg']);workbook.save(path);workbook.close()
    first=load_student_mapping(path)
    assert dict(first)=={'S123':{'name':'Example Name','face_file':'does-not-exist.jpg'}}
    path.unlink()
    assert load_student_mapping(path) is first


@pytest.mark.parametrize('rows',[[['name'],['Example']], [['rollno','name'],['S123','One'],['S123','Two']], [['rollno','name'],[None,'One']]])
def test_invalid_mapping_is_rejected(tmp_path,rows):
    workbook=Workbook()
    for row in rows: workbook.active.append(row)
    path=tmp_path/'invalid.xlsx';workbook.save(path);workbook.close()
    with pytest.raises(ValueError):load_student_mapping(path)


def test_real_opencv_decode_and_correct_name_lookup():
    encoded=cv2.QRCodeEncoder_create().encode('S123')
    encoded=cv2.copyMakeBorder(encoded,4,4,4,4,cv2.BORDER_CONSTANT,value=255)
    qr=cv2.resize(encoded,None,fx=8,fy=8,interpolation=cv2.INTER_NEAREST)
    frame=np.full((500,600,3),255,dtype=np.uint8)
    frame[140:140+qr.shape[0],80:80+qr.shape[1]]=cv2.cvtColor(qr,cv2.COLOR_GRAY2BGR)
    tracker=QRIdentityTracker({'S123':'Example Name'})
    tracks=tracker.update(frame,[person(12,(10,50,400,480))],now=100)
    identity=tracks[0]['identity']
    assert identity['track_id']==12 and identity['roll_no']=='S123' and identity['name']=='Example Name'


def test_single_card_fallback_when_multi_decode_misses():
    tracker=scanner()
    tracker.decoder.detectAndDecodeMulti.return_value=(False,(),None,None)
    tracker.decoder.detectAndDecode.return_value=('S123',corners(100,200),None)
    assert tracker.update(np.zeros((400,600,3),np.uint8),[person()],now=100)[0]['identity']['roll_no']=='S123'


def test_trimmed_code_and_empty_multi_corners_fall_back():
    tracker=scanner()
    tracker.decoder.detectAndDecodeMulti.return_value=(True,('',),np.array([corners(100,200)]),None)
    tracker.decoder.detectAndDecode.return_value=('  S123\n',corners(100,200),None)
    assert tracker.update(np.zeros((400,600,3),np.uint8),[person()],now=100)[0]['identity']['roll_no']=='S123'


def test_roi_upscale_corners_are_mapped_back_and_associated_globally():
    tracker=QRIdentityTracker({'S123':'Example'},association_margin_px=40)
    decoder=Mock();tracker.decoder=decoder
    def decode(image):
        # Person box [100,100,300,350] padded to [60,60,340,390], enlarged 3x.
        if image.shape[:2]==(990,840):
            return True,('S123',),np.array([corners(420,420)]),None
        return False,(),None,None
    decoder.detectAndDecodeMulti.side_effect=decode
    decoder.detectAndDecode.return_value=('',None,None)
    output=tracker.update(np.zeros((480,640,3),np.uint8),[person(24,(100,100,300,350))],now=100)
    assert output[0]['identity']['track_id']==24
    assert any(np.allclose(points.mean(axis=0),(200,200)) for _,points in tracker.debug_detections)


def test_explicit_pixel_margin_and_rejection_reasons():
    tracker=QRIdentityTracker({'S123':'Example'},max_distance_fraction=0,association_margin_px=40)
    details={}
    assert tracker.associate(corners(285,200),[person()],(400,600),details)==12
    assert details['reason']=='nearest_box_within_margin'
    assert tracker.associate(corners(295,200),[person()],(400,600),details) is None
    assert details['reason']=='qr_too_far'
    assert tracker.associate(corners(275,200),[person(),person(7,(300,20,500,350))],(400,600),details) is None
    assert details['reason']=='ambiguous_nearby_people'


def test_debug_overlay_is_optional_and_no_text_detection_keeps_corners():
    tracker=scanner(values=('',),debug_overlay=True)
    frame=np.zeros((400,600,3),np.uint8)
    tracker.update(frame,[person()],now=100,frame_number=33)
    assert tracker.debug_detections
    before=frame.copy();tracker.draw_debug(frame)
    assert np.any(before!=frame)
    tracker.debug_overlay=False
    before=frame.copy();tracker.draw_debug(frame)
    assert np.array_equal(before,frame)


def test_isolated_worker_decodes_without_blocking_camera_updates():
    import time
    encoded=cv2.QRCodeEncoder_create().encode('S123')
    encoded=cv2.copyMakeBorder(encoded,4,4,4,4,cv2.BORDER_CONSTANT,value=255)
    qr=cv2.resize(encoded,None,fx=8,fy=8,interpolation=cv2.INTER_NEAREST)
    frame=np.full((400,600,3),255,np.uint8)
    frame[100:100+qr.shape[0],50:50+qr.shape[1]]=cv2.cvtColor(qr,cv2.COLOR_GRAY2BGR)
    tracker=QRIdentityTracker({'S123':'Example'},isolate_decoder=True,scan_interval=.25,scan_timeout=5)
    try:
        deadline=time.monotonic()+8
        while time.monotonic()<deadline:
            tracks=tracker.update(frame,[person(24,(0,0,400,390))],frame_number=30)
            if tracks[0]['identity'] is not None:
                assert tracks[0]['identity']['roll_no']=='S123'
                break
            time.sleep(.03)
        else:pytest.fail('QR worker did not return a match')
    finally:tracker.close()


def test_decoder_timeout_terminates_worker_and_keeps_unknown(monkeypatch):
    import queue
    import ai_engine.identity.qr_identity as module
    tracker=QRIdentityTracker({'S123':'Example'},isolate_decoder=True,scan_timeout=2)
    tracker._pending=True;tracker._job_started=0;tracker._job_frame=99
    tracker._connection=Mock();tracker._connection.get_nowait.side_effect=queue.Empty
    worker=Mock();worker.is_alive.side_effect=[True,False];tracker._worker=worker
    monkeypatch.setattr(module.time,'monotonic',lambda:100)
    result=tracker.update(np.zeros((400,600,3),np.uint8),[person()],now=100)
    worker.terminate.assert_called_once()
    assert tracker._worker is None and result[0]['identity'] is None


def test_delayed_worker_result_cannot_reassociate_to_another_person(monkeypatch):
    tracker=QRIdentityTracker({'S123':'Example'},isolate_decoder=True)
    info=dict(selected_track=7,known_student=True)
    monkeypatch.setattr(tracker,'_poll_or_submit',lambda *args:[('S123',corners(100,200),info)])
    result=tracker.update(np.zeros((400,600,3),np.uint8),[person(24)],now=100)
    assert result[0]['identity'] is None


def test_partial_match_is_used_while_other_rois_are_still_decoding(monkeypatch):
    import ai_engine.identity.qr_identity as module
    tracker=QRIdentityTracker({'S123':'Example'},isolate_decoder=True)
    tracker._pending=True;tracker._job_started=99
    tracker._connection=Mock()
    info=dict(selected_track=12,known_student=True)
    tracker._connection.get_nowait.return_value=('partial',[('S123',corners(100,200),info)])
    monkeypatch.setattr(module.time,'monotonic',lambda:100)
    result=tracker.update(np.zeros((400,600,3),np.uint8),[person()],now=100)
    assert tracker._pending and result[0]['identity']['roll_no']=='S123'


@pytest.mark.parametrize('value',['','not-enrolled','https://example.com'])
def test_unknown_and_invalid_codes_are_ignored(value):
    tracker=scanner(values=(value,))
    assert tracker.update(np.zeros((400,600,3),np.uint8),[person()],now=100)[0]['identity'] is None


def test_correct_track_and_ambiguous_overlapping_or_between_people():
    tracker=scanner()
    tracks=[person(7,(300,10,500,350)),person(12)]
    output=tracker.update(np.zeros((400,600,3),np.uint8),tracks,now=100)
    assert output[0]['identity'] is None and output[1]['identity']['track_id']==12
    assert tracker.associate(corners(200,200),[person(1),person(2,(100,10,350,350))],(400,600)) is None
    assert tracker.associate(corners(275,200),tracks,(400,600)) is None
    assert tracker.associate(corners(580,390),tracks,(400,600)) is None
    assert tracker.associate(corners(255,200),[person()],(400,600))==12


def test_timeout_disappearance_stale_tracks_refresh_and_throttle():
    frame=np.zeros((400,600,3),np.uint8);tracker=scanner(timeout=30,scan_interval=1)
    assert tracker.update(frame,[person()],now=100)[0]['identity']
    tracker.update(frame,[person()],now=100.5)
    assert tracker.decoder.detectAndDecodeMulti.call_count==1
    tracker.update(frame,[person()],now=120)
    tracker.decoder.detectAndDecodeMulti.return_value=(False,(),None,None)
    assert tracker.update(frame,[person()],now=149)[0]['identity']
    assert tracker.update(frame,[person()],now=150)[0]['identity'] is not None
    tracker=scanner();tracker.update(frame,[person()],now=100);tracker.update(frame,[],now=101)
    tracker.decoder.detectAndDecodeMulti.return_value=(False,(),None,None)
    assert tracker.update(frame,[person()],now=102)[0]['identity'] is not None
    tracker=scanner();tracker.update(frame,[person()],now=100)
    assert tracker.update(frame,[person(time_since_update=1)],now=101)[0]['identity'] is not None


def test_decoder_failure_and_conflicting_codes_keep_unknown():
    frame=np.zeros((400,600,3),np.uint8);tracker=scanner()
    tracker.decoder.detectAndDecodeMulti.side_effect=cv2.error('Unreadable QR')
    assert tracker.update(frame,[person()],now=100)[0]['identity'] is None
    tracker=scanner(values=('S123','S456'),points=[corners(100,200),corners(130,200)])
    tracker.mapping={'S123':'One','S456':'Two'}
    assert tracker.update(frame,[person()],now=100)[0]['identity'] is None
    tracker=scanner(values=('S123','S123'),points=[corners(100,200),corners(350,200)])
    assert all(t['identity'] is None for t in tracker.update(frame,[person(1),person(2,(300,10,500,350))],now=100))


@pytest.mark.parametrize('activity',['sleeping','talking','running'])
def test_identity_activity_overlay_preserves_exact_track_and_unknown_behavior(activity,monkeypatch):
    frame=np.zeros((400,600,3),np.uint8);tracker=scanner()
    tracks=tracker.update(frame,[person(),person(7,(300,20,500,350))],now=100)
    for track in tracks:
        track['behaviors']=[dict(type=activity,severity='low',track_id=track['track_id'],detected_at=100)]
    texts=[]
    monkeypatch.setattr(cv2,'putText',lambda frame,text,*args,**kwargs:texts.append(text))
    FrameAnnotator().draw(frame,tracks)
    assert 'S123 | EXAMPLE NAME' in texts and 'UNKNOWN' in texts
    assert texts.count(activity.upper())==2
    assert texts.index('S123 | EXAMPLE NAME')<texts.index('#12 PERSON 0.91')


@pytest.mark.parametrize('camera_name', ['General Entrance', 'Classroom'])
def test_opt_in_identity_lifetime_and_alerts(camera_name, monkeypatch):
    from ai_engine.vision.confirmed_alerts import ConfirmedAlertTransitions
    frame = np.zeros((400,600,3),np.uint8)
    tracker = scanner(track_grace_secs=3)
    tracker.mapping = {'S123': {'name':'Example Name', 'face_file':'example.jpg'}}
    tracks = tracker.update(frame,[person()],now=100)
    identity = tracks[0]['identity']
    assert identity['face_file']=='example.jpg' and identity['identified_at']==identity['verified_at']
    tracker.decoder.detectAndDecodeMulti.return_value=(False,(),None,None)
    # Neither QR absence nor the old 30-second timeout clears an active track.
    for now in (101,140,400,1000):
        tracks = tracker.update(frame,[person()],now=now)
        assert tracks[0]['identity']['identified_at']==identity['identified_at']
    events=[{'type':'phone_use','severity':'low','track_id':12}]
    tracks[0]['behaviors']=[{**events[0],'detected_at':1000},
                           {'type':'talking','severity':'low','track_id':12,'detected_at':999},
                           {'type':'sleeping','severity':'low','track_id':12,'detected_at':998}]
    texts=[]
    monkeypatch.setattr(cv2,'putText',lambda frame,text,*args,**kwargs:texts.append(text))
    FrameAnnotator().draw(frame,tracks)
    assert 'S123 | EXAMPLE NAME' in texts and 'PHONE USE | TALKING' in texts
    assert 'SLEEPING' not in texts and not any('DEBUG' in text for text in texts)
    router=ConfirmedAlertTransitions()
    alert=router.update(events,tracks,now=1000)[0]
    assert alert['title']=='PHONE USE — S123 Example Name'
    assert alert['metadata']['student_identity']['roll_no']=='S123'
    assert router.update(events,tracks,now=1001)==[]
    tracker.update(frame,[],now=1001)
    assert tracker.update(frame,[person()],now=1002)[0]['identity'] is not None
    tracker.update(frame,[],now=1003)
    tracker.update(frame,[],now=1006)
    assert tracker.update(frame,[person()],now=1007)[0]['identity'] is None
    unknown=ConfirmedAlertTransitions().update(events,[person()],now=1007)[0]
    assert unknown['title']=='PHONE USE — Person #12'
    assert 'student_identity' not in unknown['metadata']


def test_new_qr_only_replaces_its_associated_track_and_owner_is_protected():
    frame=np.zeros((400,600,3),np.uint8)
    tracker=scanner()
    tracks=[person(),person(7,(300,20,500,350))]
    tracker.update(frame,tracks,now=100)
    tracker.mapping={'S123':'One','S456':'Two'}
    tracker.decoder.detectAndDecodeMulti.return_value=(True,('S456',),np.array([corners(350,200)]),None)
    output=tracker.update(frame,tracks,now=101)
    assert [t['identity']['roll_no'] for t in output]==['S123','S456']
    tracker.decoder.detectAndDecodeMulti.return_value=(True,('S123',),np.array([corners(350,200)]),None)
    output=tracker.update(frame,tracks,now=102)
    assert [t['identity']['roll_no'] for t in output]==['S123','S456']
    tracker.decoder.detectAndDecodeMulti.return_value=(True,('S456',),np.array([corners(100,200)]),None)
    # Expire previous owner before deliberate reassignment to this track.
    tracker.update(frame,[person()],now=103)
    output=tracker.update(frame,[person()],now=107)
    assert output[0]['identity']['roll_no']=='S456'


def test_explicit_reset_and_new_camera_session_clear_identity():
    frame=np.zeros((400,600,3),np.uint8)
    tracker=scanner()
    tracker.update(frame,[person()],now=100)
    tracker.decoder.detectAndDecodeMulti.return_value=(False,(),None,None)
    tracker.reset(12)
    assert tracker.update(frame,[person()],now=101)[0]['identity'] is None
    replacement=scanner(values=('',))
    assert replacement.update(frame,[person()],now=102)[0]['identity'] is None


def test_new_valid_qr_replaces_same_track_and_refresh_does_not_change_identified_at():
    frame=np.zeros((400,600,3),np.uint8);tracker=scanner()
    first=tracker.update(frame,[person()],now=100)[0]['identity']
    second=tracker.update(frame,[person()],now=101)[0]['identity']
    assert first['identified_at']==second['identified_at']
    tracker.mapping={'S123':'One','S456':'Two'}
    tracker.decoder.detectAndDecodeMulti.return_value=(True,('S456',),np.array([corners(100,200)]),None)
    assert tracker.update(frame,[person()],now=102)[0]['identity']['roll_no']=='S456'


@pytest.mark.parametrize('unsafe', ['../secret.jpg', 'C:\\private\\photo.jpg', '/etc/private'])
def test_profile_metadata_rejects_paths_outside_profile_folder(tmp_path,unsafe):
    workbook=Workbook();workbook.active.append(['rollno','name','face_file'])
    workbook.active.append(['S123','Example',unsafe])
    path=tmp_path/'students.xlsx';workbook.save(path);workbook.close()
    assert load_student_mapping(path)['S123']['face_file'] is None
