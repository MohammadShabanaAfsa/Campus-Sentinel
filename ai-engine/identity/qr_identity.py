"""Camera-thread-owned QR identity cache using the existing person tracker."""
import math
import time
import threading
import json
import queue
from pathlib import Path
from types import MappingProxyType

import cv2
import numpy as np
from loguru import logger

_mappings = {}
_mapping_lock = threading.Lock()


def _scan_worker(jobs, results, mapping, options):
    """QR-only child: no camera capture, tracker, DB, or student-photo access."""
    scanner = QRIdentityTracker(mapping, **options)
    try:
        while True:
            frame, tracks, frame_number = jobs.get()
            try:
                def progress(items):
                    if any(info.get('known_student') and info.get('selected_track') is not None for _,_,info in items):
                        results.put(('partial',items))
                observations=scanner._scan(frame,tracks,frame_number,time.monotonic(),on_result=progress)
                results.put(('complete',observations))
            except Exception:
                results.put(('complete',[]))
    except (EOFError, OSError):
        pass
    finally:
        jobs.close()
        results.close()


def load_student_mapping(path):
    """Cache enrollment metadata once; photos are never used for recognition."""
    path = Path(path).resolve()
    with _mapping_lock:
        if path not in _mappings:
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = book.active.iter_rows(values_only=True)
                header = [str(v).strip().lower() if v is not None else "" for v in next(rows, ())]
                if any(header.count(key) != 1 for key in ("rollno", "name")):
                    raise ValueError("Student workbook requires unique rollno and name columns")
                ri, ni = header.index("rollno"), header.index("name")
                fi = header.index('face_file') if header.count('face_file') == 1 else None
                mapping = {}
                for row in rows:
                    roll = str(row[ri]).strip() if row[ri] is not None else ""
                    name = str(row[ni]).strip() if row[ni] is not None else ""
                    if not roll and not name:
                        continue
                    if not roll or not name or roll in mapping:
                        raise ValueError("Student workbook has empty metadata or duplicate roll numbers")
                    face_file = str(row[fi]).strip() if fi is not None and row[fi] is not None else None
                    # Only a relative profile reference; never expose arbitrary paths.
                    if face_file and (Path(face_file).is_absolute() or face_file.startswith(('/', '\\')) or ':' in face_file or
                                      '..' in face_file.replace('\\', '/').split('/')):
                        face_file = None
                    mapping[roll] = {'name': name, 'face_file': face_file}
                _mappings[path] = MappingProxyType(mapping)
                logger.info(f"Loaded {len(mapping)} QR student identities.")
            finally:
                book.close()
        return _mappings[path]


class QRIdentityTracker:
    def __init__(self, mapping, timeout=30.0, scan_interval=1.0,
                 max_distance_fraction=0.08, ambiguity_margin_fraction=0.03,
                 association_margin_px=40, debug_logging=False, debug_overlay=False,
                 isolate_decoder=False, scan_timeout=2.0, track_grace_secs=3.0):
        self.mapping = mapping
        # Legacy timeout accepted for compatibility; QR absence never expires identity.
        self.scan_interval = max(0.1, scan_interval)
        self.max_distance = max(0.0, max_distance_fraction)
        self.ambiguity_margin = max(0.0, ambiguity_margin_fraction)
        self.association_margin = max(0.0, association_margin_px)
        self.debug_logging = debug_logging
        self.debug_overlay = debug_overlay
        self.debug_detections = []
        self._debug_until = 0.0
        self.isolate_decoder = isolate_decoder
        self.scan_timeout = max(.5,scan_timeout)
        self._worker = None
        self._connection = None
        self._jobs = None
        self._pending = False
        self._job_started = 0.0
        self._job_frame = None
        self.decoder = cv2.QRCodeDetector()
        self._last_seen = {}
        self._missing_since = {}
        self.track_grace_secs = max(0.0, track_grace_secs)
        self._identities = {}
        self._next_scan = 0.0
        self._last_warning = -math.inf

    def associate(self, corners, tracks, shape, details=None):
        details = details if details is not None else {}
        details.update(candidate_tracks=[], selected_track=None, reason="invalid_corners")
        points = np.asarray(corners, dtype=float).reshape(-1, 2)
        if len(points) != 4 or not np.isfinite(points).all():
            return None
        x, y = points.mean(axis=0)
        details["qr_center"] = [round(float(x),1), round(float(y),1)]
        h, w = shape[:2]
        if not (0 <= x < w and 0 <= y < h):
            details["reason"] = "qr_center_outside_frame"
            return None
        candidates, containing = [], []
        for track in tracks:
            if track.get("class_name") != "person" or track.get("time_since_update", 0) > 0:
                continue
            x1, y1, x2, y2 = track["bbox"]
            if not all(math.isfinite(v) for v in (x1,y1,x2,y2)) or x2 <= x1 or y2 <= y1:
                continue
            distance = math.hypot(max(x1-x, 0, x-x2), max(y1-y, 0, y-y2))
            if distance == 0:
                containing.append(track["track_id"])
            candidates.append((distance, track["track_id"]))
        details["candidate_tracks"] = [tid for _,tid in candidates]
        if containing:
            selected = containing[0] if len(containing) == 1 else None
            details.update(selected_track=selected, reason="containing_box" if selected is not None else "ambiguous_overlapping_boxes")
            return selected
        candidates.sort()
        diagonal = math.hypot(w, h)
        if not candidates or candidates[0][0] > max(self.association_margin, self.max_distance * diagonal):
            details["reason"] = "no_current_person" if not candidates else "qr_too_far"
            return None
        if len(candidates) > 1 and candidates[1][0] - candidates[0][0] <= self.ambiguity_margin * diagonal:
            details["reason"] = "ambiguous_nearby_people"
            return None
        details.update(selected_track=candidates[0][1], reason="nearest_box_within_margin")
        return candidates[0][1]

    def _decode(self, image, source, offset, scale, frame_shape, tracks, frame_number, now):
        observations = []
        for mode in ("multi", "single"):
            try:
                if mode == "multi":
                    found, values, corners, _ = self.decoder.detectAndDecodeMulti(image)
                    # Keep corners even when no text was decoded.
                    pairs = list(zip(values or (), corners)) if corners is not None else []
                    if corners is not None and not pairs:
                        pairs = [("", p) for p in corners]
                else:
                    value, corners, _ = self.decoder.detectAndDecode(image)
                    found = corners is not None
                    pairs = [(value, corners)] if corners is not None else []
                if not pairs:
                    pairs = [("", None)]
                for value, points in pairs:
                    value = value.strip() if isinstance(value,str) else ""
                    points = (np.asarray(points,dtype=float).reshape(4,2) / scale + np.asarray(offset)) if points is not None else None
                    info = dict(frame_number=frame_number, scan_time=round(time.time(),3),
                                source=source, decoder=mode, found=bool(found), decoded=value,
                                corners=points.round(1).tolist() if points is not None else None,
                                known_student=value in self.mapping,
                                candidate_tracks=[t['track_id'] for t in tracks if t.get('class_name')=='person' and t.get('time_since_update',0)==0], selected_track=None,
                                reason="no_qr_decoded" if not value else "unknown_student")
                    if points is not None:
                        association = {}
                        selected = self.associate(points,tracks,frame_shape,association)
                        info.update(association)
                        if not value:
                            info["selected_track"] = None
                            info["reason"] = "qr_detected_but_text_empty"
                        elif value not in self.mapping:
                            info.update(selected_track=None, reason="unknown_student")
                        elif selected is not None:
                            info["reason"] = "association_candidate"
                    if self.debug_logging:
                        logger.info("QR scan: " + json.dumps(info, ensure_ascii=True))
                    observations.append((value, points, info))
                # Single fallback also runs if any multi-detected code was empty.
                if mode == "multi" and pairs and all(isinstance(v,str) and v.strip() for v,_ in pairs):
                    break
            except Exception as error:
                if self.debug_logging:
                    logger.info(f"QR scan: frame={frame_number} source={source} decoder={mode} reason=decoder_error error={type(error).__name__}")
        return observations

    def _scan(self, frame, tracks, frame_number, now, on_result=None):
        def decode(image,source,offset,scale):
            items=self._decode(image,source,offset,scale,frame.shape,tracks,frame_number,now)
            if on_result is not None:
                on_result(items)
            return items
        observations = decode(frame,"full_frame",(0,0),1)
        def assigned_ids():
            return {info["selected_track"] for _,_,info in observations if info["known_student"] and info["selected_track"] is not None}
        if not assigned_ids():
            scale = min(2.0,1600 / frame.shape[1])
            if scale > 1:
                enlarged=cv2.resize(frame,None,fx=scale,fy=scale,interpolation=cv2.INTER_NEAREST)
                observations.extend(decode(enlarged,"upscaled_frame",(0,0),scale))
        for track in tracks:
            if track.get("class_name") != "person" or track.get("time_since_update",0)>0 or track["track_id"] in assigned_ids():
                continue
            x1,y1,x2,y2=track["bbox"]
            if x2-x1 < 64 or y2-y1 < 96:
                continue
            margin=int(self.association_margin)
            x1,y1=max(0,int(x1)-margin),max(0,int(y1)-margin)
            x2,y2=min(frame.shape[1],int(x2)+margin),min(frame.shape[0],int(y2)+margin)
            roi=frame[y1:y2,x1:x2]
            if not roi.size:
                continue
            observations.extend(decode(roi,f"person_roi_{track['track_id']}",(x1,y1),1))
            if track["track_id"] not in assigned_ids():
                scale=min(3.0,1200/max(roi.shape[:2]))
                if scale > 1:
                    enlarged=cv2.resize(roi,None,fx=scale,fy=scale,interpolation=cv2.INTER_NEAREST)
                    observations.extend(decode(enlarged,f"upscaled_roi_{track['track_id']}",(x1,y1),scale))
        return observations

    def draw_debug(self, frame):
        if not self.debug_overlay or time.monotonic() > self._debug_until:
            return frame
        for value,points in self.debug_detections:
            points=np.rint(points).astype(np.int32)
            cv2.polylines(frame,[points],True,(255,220,0),2)
            x,y=points.min(axis=0)
            cv2.putText(frame,"QR: "+(value or "UNREADABLE"),(max(0,int(x)),max(18,int(y)-5)),
                        cv2.FONT_HERSHEY_SIMPLEX,.5,(255,220,0),1,cv2.LINE_AA)
        return frame

    def close(self):
        if self._connection is not None:
            self._connection.cancel_join_thread()
            self._connection.close()
            self._connection = None
        if self._jobs is not None:
            self._jobs.cancel_join_thread()
            self._jobs.close()
            self._jobs = None
        if self._worker is not None:
            if self._worker.is_alive():
                self._worker.terminate()
            self._worker.join(timeout=1)
            if self._worker.is_alive():
                self._worker.kill()
                self._worker.join(timeout=1)
            self._worker = None
        self._pending = False

    def reset(self, track_id=None):
        """Explicit reset also cancels old decoder results before clearing identity."""
        self.close()
        for state in (self._identities, self._last_seen, self._missing_since):
            if track_id is None:
                state.clear()
            else:
                state.pop(track_id, None)
        self.debug_detections = []
        self._next_scan = 0.0

    def _poll_or_submit(self, frame, tracks, frame_number, now, active):
        observations = None
        if self._pending:
            try:
                try:
                    kind,observations=self._connection.get_nowait()
                    self._pending=kind!='complete'
                    if self.debug_logging:
                        for _,_,info in observations:
                            logger.info("QR scan: " + json.dumps(info,ensure_ascii=True))
                except queue.Empty:
                    pass
                if self._pending and time.monotonic()-self._job_started > self.scan_timeout:
                    if self.debug_logging:
                        logger.info(f"QR scan: frame={self._job_frame} decoded='' selected_track=None reason=decoder_timeout timeout_secs={self.scan_timeout}")
                    self.close()
                    self._next_scan=now+1  # Back off after a pathological native decode.
            except (EOFError,OSError):
                self.close()
        if active and not self._pending and now>=self._next_scan:
            self._next_scan=now+self.scan_interval
            try:
                if self._worker is None:
                    import multiprocessing
                    context=multiprocessing.get_context('spawn')
                    self._jobs=context.Queue(maxsize=1)
                    self._connection=context.Queue(maxsize=1)
                    options=dict(association_margin_px=self.association_margin,
                                 max_distance_fraction=self.max_distance,
                                 ambiguity_margin_fraction=self.ambiguity_margin,
                                 debug_logging=False)  # Log worker records in the parent file.
                    self._worker=context.Process(target=_scan_worker,args=(self._jobs,self._connection,dict(self.mapping),options),daemon=True)
                    self._worker.start()
                self._jobs.put_nowait((frame.copy(),[{**t,'bbox':list(t['bbox'])} for t in tracks],frame_number))
                self._pending=True
                self._job_started=time.monotonic()
                self._job_frame=frame_number
                if self.debug_logging:
                    logger.info(f"QR scan submitted: frame={frame_number} time={time.time():.3f} candidate_tracks={sorted(active)}")
            except Exception:
                self.close()
        return observations

    def update(self, frame, tracks, now=None, frame_number=None):
        now = time.monotonic() if now is None else now
        active = {t["track_id"] for t in tracks if t.get("class_name") == "person" and t.get("time_since_update", 0) == 0}
        for tid in self._identities:
            if tid not in active:self._missing_since.setdefault(tid,self._last_seen.get(tid,now))
        self._identities = {tid: value for tid, value in self._identities.items()
                            if tid not in self._missing_since or now-self._missing_since[tid]<=self.track_grace_secs}
        for tid in active:
            self._last_seen[tid]=now;self._missing_since.pop(tid,None)
            if tid in self._identities:
                self._identities[tid]['last_seen'] = time.time()
        self._missing_since={tid:seen for tid,seen in self._missing_since.items() if tid in self._identities}
        self._last_seen={tid:seen for tid,seen in self._last_seen.items() if tid in active or now-seen<=self.track_grace_secs}
        observations = None
        if self.isolate_decoder:
            observations=self._poll_or_submit(frame,tracks,frame_number,now,active)
        elif active and now >= self._next_scan:
            self._next_scan = now + self.scan_interval
            observations=self._scan(frame, tracks, frame_number, now)
        if observations is not None:
            try:
                self.debug_detections=[(value,points) for value,points,_ in observations if points is not None]
                self._debug_until=time.monotonic()+max(.5,self.scan_interval*2)
                assignments = {}
                if observations:
                    for roll, points, info in observations:
                        if not isinstance(roll, str) or roll not in self.mapping:
                            continue
                        # Async results use the original frame's exact track ID,
                        # never re-match old QR coordinates to newer boxes.
                        tid = info.get('selected_track')
                        if tid is not None and tid in active:
                            assignments.setdefault(tid, set()).add(roll)
                # Reject conflicting codes on one person, and the same student
                # simultaneously presented by multiple people.
                owners = {}
                for tid, rolls in assignments.items():
                    for roll in rolls:
                        owners.setdefault(roll, set()).add(tid)
                for tid, rolls in assignments.items():
                    if len(rolls) != 1 or len(owners[next(iter(rolls))]) != 1:
                        if self.debug_logging:
                            logger.info(f"QR assignment: frame={frame_number} selected_track={tid} reason=conflicting_codes_or_multiple_holders")
                        continue
                    roll = next(iter(rolls))
                    if any(previous!=tid and value['roll_no']==roll for previous,value in self._identities.items()):
                        continue  # A visible/grace-protected owner cannot be transferred.
                    for previous, value in list(self._identities.items()):
                        if previous != tid and value["roll_no"] == roll:
                            del self._identities[previous]
                    student = self.mapping[roll]
                    student = {'name': student, 'face_file': None} if isinstance(student, str) else student
                    previous = self._identities.get(tid)
                    identified_at = (previous['identified_at'] if previous and previous['roll_no'] == roll else time.time())
                    self._identities[tid] = {"track_id": tid, "roll_no": roll, "name": student['name'],
                                             "face_file": student.get('face_file'),
                                             "verified": True, "method": "qr", "verified_at": identified_at,
                                             "identified_at": identified_at, "last_seen": time.time()}
                    if self.debug_logging:
                        logger.info(f"QR assignment: frame={frame_number} selected_track={tid} reason=assigned")
            except Exception as error:
                if now - self._last_warning >= 30:
                    logger.warning(f"QR scan unavailable ({type(error).__name__}); video and behavior analysis continue")
                    self._last_warning = now
        return [{**track, "identity": dict(self._identities[track["track_id"]])
                 if track.get("class_name") == "person" and track["track_id"] in self._identities else None}
                for track in tracks]
